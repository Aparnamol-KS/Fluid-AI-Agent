import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = PROJECT_ROOT / "outputs"

BLUE = "1F4E79"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "F3F7FB"
DARK_TEXT = "2D3748"
MUTED_TEXT = "64748B"
CODE_BACKGROUND = "F3F4F6"
BORDER_COLOR = "D6DEE8"
WHITE = "FFFFFF"


def clean_inline(text: str) -> str:
    """Remove common Markdown inline syntax while preserving its readable text."""
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", "", text)
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"__(.*?)__", r"\1", text)
    text = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"\1", text)
    text = re.sub(r"(?<!_)_([^_\n]+)_(?!_)", r"\1", text)
    text = re.sub(r"~~(.*?)~~", r"\1", text)
    return text.strip()


def normalize_heading(text: str) -> str:
    text = clean_inline(text).strip().rstrip("#").strip()
    return re.sub(r"[\W_]+", "", text).casefold()


def is_table_line(line: str) -> bool:
    return "|" in line and line.count("|") >= 2


def parse_table_row(line: str) -> list[str]:
    """Split a Markdown table row while preserving escaped pipes and code spans."""
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]

    cells = []
    current = []
    escaped = False
    in_code = False
    i = 0
    while i < len(line):
        char = line[i]
        if char == "\\" and not escaped:
            escaped = True
            current.append(char)
            i += 1
            continue
        if char == "`" and not escaped:
            in_code = not in_code
            current.append(char)
        elif char == "|" and not escaped and not in_code:
            cells.append(clean_inline("".join(current).replace(r"\|", "|")))
            current = []
        else:
            current.append(char)
        escaped = False
        i += 1

    cells.append(clean_inline("".join(current).replace(r"\|", "|")))
    return cells


def is_separator_row(line: str) -> bool:
    cells = parse_table_row(line)
    return bool(cells) and all(
        re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells
    )


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)

    for edge, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = properties.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        properties.append(header)
    header.set(qn("w:val"), "true")


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def set_table_borders(table, color=BORDER_COLOR, size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_inline_runs(
    paragraph,
    text: str,
    *,
    font_name="Aptos",
    font_size=10,
    color=DARK_TEXT,
    bold=False,
    italic=False,
):
    """Add basic Markdown inline formatting without treating text as HTML."""
    pattern = re.compile(
        r"(\*\*.+?\*\*|__.+?__|(?<!\*)\*[^*\n]+\*(?!\*)|"
        r"(?<!_)_[^_\n]+_(?!_)|`[^`]+`|~~.+?~~)"
    )
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            _style_run(run, font_name, font_size, color, bold, italic)
        token = match.group(0)
        if token.startswith(("**", "__")):
            run = paragraph.add_run(token[2:-2])
            _style_run(run, font_name, font_size, color, True, italic)
        elif token.startswith("~~"):
            run = paragraph.add_run(token[2:-2])
            _style_run(run, font_name, font_size, color, bold, italic)
            run.font.strike = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            _style_run(run, "Consolas", max(font_size - 0.5, 8), color, bold, italic)
            run.font.highlight_color = 7
        else:
            run = paragraph.add_run(token[1:-1])
            _style_run(run, font_name, font_size, color, bold, True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        _style_run(run, font_name, font_size, color, bold, italic)


def _style_run(run, font_name, font_size, color, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_code_block(doc: Document, code_lines: list[str], language: str = "") -> None:
    """Render fenced code and diagrams in a shaded monospaced panel."""
    code = "\n".join(code_lines).rstrip()
    if not code:
        return

    # A one-cell table creates a stable shaded panel with consistent padding.
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = (
        doc.sections[0].page_width
        - doc.sections[0].left_margin
        - doc.sections[0].right_margin
    )
    set_cell_shading(cell, CODE_BACKGROUND)
    set_cell_margins(cell, top=120, start=150, bottom=120, end=150)
    set_table_borders(table, color=BORDER_COLOR, size="3")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True

    lines = code.split("\n")
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line.expandtabs(4))
        run.font.name = "Consolas"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(DARK_TEXT)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def add_markdown_table(doc: Document, rows: list[str]) -> None:
    """Convert Markdown pipe-table rows into a styled native Word table."""
    parsed_rows = [parse_table_row(row) for row in rows if not is_separator_row(row)]
    if not parsed_rows:
        return

    # Ignore malformed one-cell lines rather than creating a misleading table.
    column_count = max(len(row) for row in parsed_rows)
    if column_count < 2:
        doc.add_paragraph(clean_inline(rows[0]))
        return

    normalized_rows = [row + [""] * (column_count - len(row)) for row in parsed_rows]
    table = doc.add_table(rows=0, cols=column_count)
    table.style = "Table Grid"
    table.autofit = False
    section = doc.sections[0]
    available_width = section.page_width - section.left_margin - section.right_margin
    column_width = available_width / column_count
    set_table_borders(table)

    for row_index, values in enumerate(normalized_rows):
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells

        for col_index, cell in enumerate(cells):
            cell.width = column_width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.widow_control = True
            add_inline_runs(
                paragraph,
                values[col_index],
                font_size=8.5,
                bold=(row_index == 0),
                color=BLUE if row_index == 0 else DARK_TEXT,
            )
            if row_index == 0:
                set_cell_shading(cell, LIGHT_BLUE)
                paragraph.paragraph_format.keep_with_next = True
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)

        if row_index == 0:
            set_repeat_table_header(row)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def add_page_number(paragraph) -> None:
    paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def _is_bullet(line: str) -> bool:
    return bool(re.match(r"^\s*[-*+]\s+", line))


def _is_numbered(line: str) -> bool:
    return bool(re.match(r"^\s*\d+[.)]\s+", line))


def _set_document_properties(doc: Document, title: str, document_type: str) -> None:
    props = doc.core_properties
    props.title = title
    props.subject = document_type
    props.author = "Fluid AI Agent"


def create_docx(
    title: str, document_type: str, content: str, assumptions: list[str]
) -> Path:
    """Create a polished Word document from Markdown-like content."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    safe_title = re.sub(r"[^a-zA-Z0-9_-]+", "_", title).strip("_")
    safe_title = safe_title[:60] or "business_document"
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    output_path = OUTPUT_DIR / f"{safe_title}_{timestamp}.docx"

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(DARK_TEXT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True

    for style_name, size in (
        ("Title", 25),
        ("Heading 1", 16),
        ("Heading 2", 12),
        ("Heading 3", 10.5),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(4)

    _set_document_properties(doc, title, document_type)

    title_paragraph = doc.add_paragraph(style="Title")
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.space_after = Pt(4)
    title_paragraph.add_run(title)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle.paragraph_format.keep_with_next = True
    subtitle_run = subtitle.add_run(document_type.replace("_", " ").title())
    subtitle_run.italic = True
    subtitle_run.font.color.rgb = RGBColor.from_string(MUTED_TEXT)
    subtitle_run.font.size = Pt(10)

    if assumptions:
        doc.add_heading("Assumptions", level=1)
        for assumption in assumptions:
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, clean_inline(assumption), font_size=10)

    lines = content.splitlines()
    title_key = normalize_heading(title)
    assumptions_key = normalize_heading("Assumptions")
    i = 0
    in_code = False
    code_lines: list[str] = []
    code_language = ""

    while i < len(lines):
        raw_line = lines[i].rstrip()
        line = raw_line.strip()

        # Fenced code blocks, including ASCII architecture diagrams.
        fence_match = re.match(r"^\s*(```+|~~~+)\s*(.*)$", raw_line)
        if fence_match:
            # fence = fence_match.group(1)
            if not in_code:
                in_code = True
                code_lines = []
                code_language = fence_match.group(2).strip()
            else:
                add_code_block(doc, code_lines, code_language)
                in_code = False
                code_lines = []
                code_language = ""
            i += 1
            continue

        if in_code:
            code_lines.append(raw_line.expandtabs(4))
            i += 1
            continue

        if not line:
            i += 1
            continue

        # Markdown table: require a separator row directly after the header.
        if (
            is_table_line(line)
            and i + 1 < len(lines)
            and is_separator_row(lines[i + 1].strip())
        ):
            table_rows = [line, lines[i + 1].strip()]
            i += 2
            while i < len(lines) and is_table_line(lines[i].strip()):
                table_rows.append(lines[i].strip())
                i += 1
            add_markdown_table(doc, table_rows)
            continue

        # Headings from Markdown levels 1–4, mapped to Word heading styles.
        heading_match = re.match(r"^(#{1,4})\s+(.+?)\s*#*\s*$", line)
        if heading_match:
            heading_text = clean_inline(heading_match.group(2))
            heading_key = normalize_heading(heading_text)
            if heading_key == title_key:
                i += 1
                continue
            if heading_key == assumptions_key and assumptions:
                i += 1
                continue
            level = min(len(heading_match.group(1)), 3)
            doc.add_heading(heading_text, level=level)
            i += 1
            continue

        # Horizontal rules.
        if re.fullmatch(r"[-*_]{3,}", line):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(5)
            p_pr = paragraph._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), BORDER_COLOR)
            borders.append(bottom)
            p_pr.append(borders)
            i += 1
            continue

        # Lists, including simple nested indentation.
        if _is_bullet(line) or _is_numbered(line):
            indent = len(raw_line) - len(raw_line.lstrip())
            style = "List Number" if _is_numbered(line) else "List Bullet"
            list_text = re.sub(r"^\s*(?:[-*+]|\d+[.)])\s+", "", raw_line)
            paragraph = doc.add_paragraph(style=style)
            paragraph.paragraph_format.left_indent = Inches(0.25 * min(indent // 2, 3))
            add_inline_runs(paragraph, list_text, font_size=10)
            i += 1
            continue

        # Skip a repeated document-type line if the model emitted one.
        if re.match(r"^\*{0,2}\s*Document Type\s*:", line, flags=re.IGNORECASE):
            i += 1
            continue

        # Paragraphs are joined until a blank line or block-level element.
        paragraph_lines = [line]
        i += 1
        while i < len(lines):
            next_raw = lines[i].rstrip()
            next_line = next_raw.strip()
            if not next_line:
                i += 1
                break
            if (
                re.match(r"^\s*(```+|~~~+)", next_raw)
                or re.match(r"^#{1,4}\s+", next_line)
                or _is_bullet(next_line)
                or _is_numbered(next_line)
                or re.fullmatch(r"[-*_]{3,}", next_line)
                or (
                    is_table_line(next_line)
                    and i + 1 < len(lines)
                    and is_separator_row(lines[i + 1].strip())
                )
            ):
                break
            paragraph_lines.append(next_line)
            i += 1

        paragraph_text = " ".join(paragraph_lines)
        paragraph = doc.add_paragraph()
        add_inline_runs(paragraph, paragraph_text, font_size=10)
        paragraph.paragraph_format.widow_control = True

    # Gracefully render a missing closing fence rather than dropping code.
    if in_code and code_lines:
        add_code_block(doc, code_lines, code_language)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(3)
    footer_run = footer.add_run("Fluid AI Agent  |  ")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(MUTED_TEXT)
    add_page_number(footer)

    doc.save(output_path)
    return output_path
